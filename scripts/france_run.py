"""
FRANCE-RUN — build the agriculteurs deliverable (producteurs + éleveurs) for
every département of métropole
==========================================================================
Sam, 2026-09-19: "lance l'extraction des agriculteurs sur toute la France dès
que possible, c'est assez urgent. Pas la peine de chercher d'autres annuaires
ou de réaliser des extractions quotidiennes, on se concentre sur l'existant."

So this driver runs the EXISTING chain, unchanged, once per département. It
adds no source. It is a loop with a status file, not a new pipeline.

    global phase (once)   Agence Bio per dept -> its adapter; the Supabase
                          provider file and contact_points claims for all of
                          France (one SELECT each, read-only)
    per département       registry pull (M6 éleveurs + M7 producteurs) ->
                          transform -> match -> export -> gate -> copy the
                          gated xlsx into exports/france_agriculture/<Région>/

Every step is the ordinary script, called as a subprocess, and every one is
idempotent: an acquire resumes from its sidecars and returns immediately when
a cell is complete, so re-running the driver after a crash, a reboot or a
CTRL-C costs nothing. That is the whole resume story — there is no separate
state to repair.

A département that fails is written to france_status.csv with the step that
failed, and THE RUN CONTINUES. One bad département must not cost the other
95. A département whose gate fails is built but NOT copied to the delivery
folder: we never ship an ungated file.

    python scripts/france_run.py --departements 15,32,2B     # a pilot
    python scripts/france_run.py --departements Occitanie    # a région
    python scripts/france_run.py --all                       # France
    python scripts/france_run.py --rebuild-only              # no network:
                                  match -> export -> gate again, to pick up
                                  harvests that finished after the first pass
    python scripts/france_run.py --status                    # where we are
    python scripts/france_run.py --recap                     # the recap xlsx,
                                  counted by reading every xlsx back

Long runs must be started DETACHED (the Claude tool kills a background job
after 10 minutes) — see scripts/france_overnight.sh.
"""

import argparse
import csv
import gzip
import logging
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))
from france_lib import (METRO_DEPARTEMENTS, REGION_OF, ALREADY_DONE,   # noqa: E402
                        parse_departements, region_dir)

PY = sys.executable
SCRIPTS = PROJECT_ROOT / "scripts"
OUT_ROOT = PROJECT_ROOT / "exports" / "france_agriculture"
STATUS_PATH = OUT_ROOT / "france_status.csv"
RECAP_PATH = OUT_ROOT / "france_recap.xlsx"
LOG_DIR = OUT_ROOT / "logs"

M7_CHECK = PROJECT_ROOT / "exports" / "producteurs" / "checkpoints"
M6_CHECK = PROJECT_ROOT / "exports" / "eleveurs" / "checkpoints"
M7_OUT = PROJECT_ROOT / "exports" / "producteurs"
M6_OUT = PROJECT_ROOT / "exports" / "eleveurs"

VERSION = "v1"
PORCINS = "01.46Z"          # Ines's principle: never pulled at all

STATUS_FIELDS = ["dept", "region", "status", "failed_step", "rows", "producteurs", "eleveurs",
                 "phones", "emails", "provider_phones", "joignables", "sans_siret",
                 "gate", "started", "finished", "minutes"]

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s %(levelname)-7s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("france")


# ---------------------------------------------------------------- running

def run(step: str, dept: str, args: list[str], timeout: int = 7200) -> tuple[bool, str]:
    """One pipeline script as a subprocess. Its full output goes to a per-dept
    log file; only the outcome comes back here, so a 96-département run stays
    readable."""
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    logfile = LOG_DIR / f"{dept}.log"
    cmd = [PY, str(SCRIPTS / f"{step}.py")] + args
    t0 = time.time()
    with logfile.open("a", encoding="utf-8") as fh:
        fh.write(f"\n{'=' * 70}\n{datetime.now():%Y-%m-%d %H:%M:%S}  {' '.join(cmd[1:])}\n{'=' * 70}\n")
        fh.flush()
        try:
            p = subprocess.run(cmd, stdout=fh, stderr=subprocess.STDOUT,
                               cwd=str(PROJECT_ROOT), timeout=timeout)
        except subprocess.TimeoutExpired:
            fh.write(f"\n*** TIMEOUT after {timeout}s ***\n")
            log.error(f"[{dept}] {step}: TIMEOUT after {timeout // 60} min")
            return False, "timeout"
    dt = time.time() - t0
    if p.returncode != 0:
        tail = ""
        try:
            tail = logfile.read_text(encoding="utf-8", errors="replace").strip().splitlines()[-1][:160]
        except Exception:
            pass
        log.error(f"[{dept}] {step}: exit {p.returncode} after {dt:.0f}s — {tail}")
        return False, f"exit {p.returncode}"
    log.info(f"[{dept}] {step}: ok ({dt:.0f}s)")
    return True, ""


def gzip_raw(dept: str) -> None:
    """The raw JSONL is ~35 MB per département and its only job is to let a
    transform be replayed. Gzip it (about 8x) once the transform has run, so
    96 départements do not leave 3 GB of uncompressed JSON in OneDrive. The
    acquire sidecars are untouched, so a re-pull still resumes correctly."""
    for d in (M7_CHECK, M6_CHECK):
        raw = d / f"api_raw_{dept}.jsonl"
        if not raw.exists() or raw.stat().st_size == 0:
            continue
        gz = raw.with_suffix(".jsonl.gz")
        try:
            with raw.open("rb") as src, gzip.open(gz, "wb", compresslevel=6) as dst:
                shutil.copyfileobj(src, dst, length=1 << 20)
            raw.unlink()
        except Exception as exc:
            log.warning(f"[{dept}] could not gzip {raw.name}: {exc}")


# ---------------------------------------------------------------- reading back

def read_xlsx_counts(dept: str) -> dict:
    """Count the deliverable by READING THE XLSX BACK, never from the logs —
    the rule this project learned the hard way (written= vs collected=)."""
    path = M7_OUT / f"producteurs_{dept}_{VERSION}.xlsx"
    if not path.exists():
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Producteurs"]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return {}
        ix = {name: i for i, name in enumerate(header)}

        def cell(r, name):
            i = ix.get(name)
            return r[i] if i is not None and i < len(r) else None

        n = tel = mail = prov = reach = n_prod = n_elev = 0
        for r in rows:
            if r is None or all(v in (None, "") for v in r):
                continue
            n += 1
            t, m = cell(r, "telephone_final"), cell(r, "email_final")
            if t:
                tel += 1
            if m:
                mail += 1
            if t or m:
                reach += 1
            if cell(r, "provider_phone"):
                prov += 1
            # population_source is the unambiguous split the export already
            # writes: 'registre' = producteurs, 'eleveurs (M6)' = livestock.
            if "eleveurs" in (cell(r, "population_source") or ""):
                n_elev += 1
            else:
                n_prod += 1
        sans = 0
        if "Sans SIRET" in wb.sheetnames:
            s = wb["Sans SIRET"].iter_rows(values_only=True)
            next(s, None)
            sans = sum(1 for r in s if r and any(v not in (None, "") for v in r))
        return {"rows": n, "producteurs": n_prod, "eleveurs": n_elev, "phones": tel,
                "emails": mail, "provider_phones": prov, "joignables": reach,
                "sans_siret": sans}
    finally:
        wb.close()


# ---------------------------------------------------------------- status file

def status_paths() -> list[Path]:
    """Every worker's status file. Workers never share one file — two
    processes rewriting the same CSV lose each other's rows."""
    return sorted(OUT_ROOT.glob("france_status*.csv")) if OUT_ROOT.exists() else []


def load_status(own: Path | None = None) -> dict:
    """Merged view of every worker's file. `own` is read last so a worker's
    own rows win for the départements it owns."""
    out: dict = {}
    for p in status_paths() + ([own] if own and own.exists() else []):
        try:
            with p.open(encoding="utf-8-sig", newline="") as fh:
                for r in csv.DictReader(fh, delimiter=";"):
                    if r.get("dept"):
                        out[r["dept"]] = r
        except Exception as exc:
            log.warning(f"could not read {p.name}: {exc}")
    return out


def save_status(rows: dict, path: Path = STATUS_PATH) -> None:
    """Written to a .tmp then replaced, so a reader never sees a half file."""
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    order = {d: i for i, d in enumerate(METRO_DEPARTEMENTS)}
    tmp = path.with_suffix(".csv.tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=STATUS_FIELDS, delimiter=";",
                           extrasaction="ignore")
        w.writeheader()
        for d in sorted(rows, key=lambda x: order.get(x, 999)):
            w.writerow(rows[d])
    tmp.replace(path)


# ---------------------------------------------------------------- the phases

def global_phase(depts: list[str], skip_supabase: bool) -> None:
    """Whole-corpus steps: they write ONE file covering every département, so
    they run once, not per département."""
    log.info("=" * 70)
    log.info(f"GLOBAL PHASE — {len(depts)} département(s)")
    log.info("=" * 70)
    spec = ",".join(depts)

    for d in depts:                       # Agence Bio: ~6 pages of 250 per dept
        run("m3ag_s1_acquire", "_global", ["--departements", d], timeout=1800)
        run("m3ag_s2_transform", "_global", ["--departements", d], timeout=900)
    run("m6_s20_agencebio", "_global", ["--departements", spec], timeout=1800)

    if skip_supabase:
        log.info("--skip-supabase: provider_agri.csv / db_claims.csv left as they are")
    else:
        run("m7_s19_supabase", "_global", ["--departements", spec], timeout=3600)
        run("m6_s19_supabase", "_global", ["--departements", spec], timeout=3600)


def build_dept(dept: str, rebuild_only: bool) -> dict:
    """One département, end to end. Returns its status row."""
    region = region_dir(dept)
    started = datetime.now()
    st = {"dept": dept, "region": region, "status": "running", "failed_step": "",
          "started": started.strftime("%Y-%m-%d %H:%M"), "finished": "", "gate": "",
          "minutes": ""}

    steps: list[tuple[str, list[str], int]] = []
    if not rebuild_only:
        steps += [
            ("m6_s1_acquire", ["--departements", dept, "--active-only",
                               "--skip-naf", PORCINS], 10800),
            ("m6_s2_transform", ["--departements", dept], 900),
            ("m7_s1_acquire", ["--departements", dept, "--active-only"], 10800),
            ("m7_s2_transform", ["--departements", dept], 900),
        ]
    steps += [
        ("m6_s8_match", ["--departement", dept], 3600),
        ("m6_s9_export", ["--departement", dept, "--version", VERSION], 3600),
        ("m6_s12_check", ["--departement", dept, "--version", VERSION, "--strict"], 1800),
        ("m7_s8_match", ["--departement", dept], 3600),
        ("m7_s9_export", ["--departement", dept, "--version", VERSION,
                          "--with-eleveurs", "--eleveurs-version", VERSION], 3600),
        ("m7_s12_check", ["--departement", dept, "--version", VERSION, "--strict"], 1800),
    ]

    for step, args, timeout in steps:
        ok, why = run(step, dept, args, timeout)
        if not ok:
            # A gate failure is a verdict, not a crash: the file exists but is
            # not shippable. Anything else is a build failure.
            st["status"] = "gate_failed" if step.endswith("_check") else "failed"
            st["gate"] = "FAILED" if step.endswith("_check") else ""
            st["failed_step"] = f"{step} ({why})"
            break
        if step == "m7_s2_transform":
            gzip_raw(dept)
    else:
        st["status"] = "ok"
        st["gate"] = "0 failed"

    st.update(read_xlsx_counts(dept))

    # Ship only what passed its gate.
    if st["status"] == "ok":
        dest = OUT_ROOT / region
        dest.mkdir(parents=True, exist_ok=True)
        for suffix in (".xlsx", ".csv", "_sans_siret.csv"):
            src = M7_OUT / f"producteurs_{dept}_{VERSION}{suffix}"
            if src.exists():
                shutil.copy2(src, dest / src.name)

    st["finished"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    st["minutes"] = f"{(datetime.now() - started).total_seconds() / 60:.1f}"
    return st


# ---------------------------------------------------------------- reporting

def print_status() -> None:
    rows = load_status()
    if not rows:
        print("no run yet (no france_status.csv)")
        return
    by = {}
    for r in rows.values():
        by.setdefault(r["status"], []).append(r["dept"])
    print(f"\n{'dept':<6}{'région':<28}{'status':<13}{'rows':>8}{'tel':>8}{'mail':>8}"
          f"{'prov.tel':>10}{'joign.':>8}  gate")
    print("-" * 97)
    order = {d: i for i, d in enumerate(METRO_DEPARTEMENTS)}
    tot = {k: 0 for k in ("rows", "phones", "emails", "provider_phones", "joignables")}
    for d in sorted(rows, key=lambda x: order.get(x, 999)):
        r = rows[d]
        print(f"{d:<6}{r['region']:<28}{r['status']:<13}{r.get('rows', ''):>8}"
              f"{r.get('phones', ''):>8}{r.get('emails', ''):>8}"
              f"{r.get('provider_phones', ''):>10}{r.get('joignables', ''):>8}  {r.get('gate', '')}"
              + (f"   <- {r['failed_step']}" if r.get("failed_step") else ""))
        for k in tot:
            try:
                tot[k] += int(r.get(k) or 0)
            except ValueError:
                pass
    print("-" * 97)
    print(f"{'TOTAL':<6}{len(rows):<28}{'':13}{tot['rows']:>8}{tot['phones']:>8}"
          f"{tot['emails']:>8}{tot['provider_phones']:>10}{tot['joignables']:>8}")
    print()
    for k, v in sorted(by.items()):
        print(f"  {k:<13} {len(v):>3}  {', '.join(sorted(v)[:30])}{' …' if len(v) > 30 else ''}")
    done = set(rows)
    left = [d for d in METRO_DEPARTEMENTS if d not in done and d not in ALREADY_DONE]
    print(f"\n  not started  {len(left):>3}  {', '.join(left[:30])}{' …' if len(left) > 30 else ''}")
    capped = M7_CHECK / "capped_cells.csv"
    if capped.exists():
        print(f"\n  ⚠ capped cells (over the API's 10,000 pagination cap): {capped}")


def write_recap() -> None:
    """One row per département, counted by reading every xlsx back."""
    from openpyxl import Workbook
    rows = load_status()
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Recap"
    cols = ["Departement", "Region", "Lignes", "Producteurs", "Eleveurs", "Telephones",
            "E-mails", "Telephone fichier client", "Joignables", "% joignables",
            "Sans SIRET", "Controle", "Statut"]
    ws.append(cols)
    order = {d: i for i, d in enumerate(METRO_DEPARTEMENTS)}
    tot = {k: 0 for k in ("rows", "producteurs", "eleveurs", "phones", "emails",
                          "provider_phones", "joignables", "sans_siret")}
    for d in sorted(rows, key=lambda x: order.get(x, 999)):
        r = rows[d]
        live = read_xlsx_counts(d) or {k: r.get(k) for k in tot}
        n = int(live.get("rows") or 0)
        for k in tot:
            try:
                tot[k] += int(live.get(k) or 0)
            except (TypeError, ValueError):
                pass
        j = int(live.get("joignables") or 0)
        ws.append([d, r["region"], n, live.get("producteurs"), live.get("eleveurs"),
                   live.get("phones"), live.get("emails"), live.get("provider_phones"),
                   j, (round(100 * j / n, 1) if n else 0), live.get("sans_siret"),
                   r.get("gate", ""), r.get("status", "")])
    ws.append([])
    ws.append(["TOTAL", f"{len(rows)} départements", tot["rows"], tot["producteurs"],
               tot["eleveurs"], tot["phones"], tot["emails"], tot["provider_phones"],
               tot["joignables"],
               (round(100 * tot["joignables"] / tot["rows"], 1) if tot["rows"] else 0),
               tot["sans_siret"], "", ""])
    wb.save(RECAP_PATH)
    log.info(f"recap -> {RECAP_PATH}")
    log.info(f"TOTAL {tot['rows']} rows · {tot['phones']} phones · {tot['emails']} e-mails · "
             f"{tot['provider_phones']} provider phones · {tot['joignables']} joignables")


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--departements", default="", help="comma list, a région name, or 'all'")
    ap.add_argument("--all", action="store_true", help="every département of métropole")
    ap.add_argument("--rebuild-only", action="store_true",
                    help="no network: match -> export -> gate again (after a harvest finished)")
    ap.add_argument("--skip-global", action="store_true", help="skip the global phase")
    ap.add_argument("--skip-supabase", action="store_true",
                    help="global phase without the two SELECTs (already pulled)")
    ap.add_argument("--redo", action="store_true",
                    help="rebuild départements already marked ok (default: skip them)")
    ap.add_argument("--include-done", action="store_true",
                    help="also rebuild 03 and 63 (their V2 files are already delivered)")
    ap.add_argument("--worker", default="",
                    help="name this worker (writes france_status_<name>.csv). Several "
                         "workers on DISJOINT département lists run in parallel; they "
                         "must never share one status file.")
    ap.add_argument("--status", action="store_true")
    ap.add_argument("--recap", action="store_true")
    args = ap.parse_args()

    if args.status:
        print_status()
        return
    if args.recap:
        write_recap()
        return

    my_path = (OUT_ROOT / f"france_status_{args.worker}.csv") if args.worker else STATUS_PATH

    depts = list(METRO_DEPARTEMENTS) if args.all else parse_departements(args.departements)
    if not depts:
        sys.exit("nothing to do: pass --departements <list|région|all> or --all")
    if not args.include_done:
        depts = [d for d in depts if d not in ALREADY_DONE]

    status = load_status(my_path)
    mine = {d: r for d, r in status.items() if d in depts}   # this worker's own file
    if not args.redo:
        skip = [d for d in depts if status.get(d, {}).get("status") == "ok"]
        if skip:
            log.info(f"already ok, skipped ({len(skip)}): {', '.join(skip[:20])}"
                     + (" …" if len(skip) > 20 else "") + "   — use --redo to rebuild")
        depts = [d for d in depts if d not in skip]
    if not depts:
        log.info("every requested département is already ok. Nothing to do.")
        return

    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    log.info(f"FRANCE RUN — {len(depts)} département(s), version {VERSION}"
             + (" (rebuild only)" if args.rebuild_only else ""))

    if not args.skip_global and not args.rebuild_only:
        global_phase(depts, args.skip_supabase)

    t0 = time.time()
    for i, dept in enumerate(depts, 1):
        log.info("-" * 70)
        log.info(f"[{i}/{len(depts)}] département {dept} ({region_dir(dept)})")
        try:
            st = build_dept(dept, args.rebuild_only)
        except KeyboardInterrupt:
            log.warning("interrupted — re-run the same command to resume")
            save_status(mine, my_path)
            raise
        except Exception as exc:                      # never let one dept kill the run
            log.exception(f"[{dept}] unexpected error: {exc}")
            st = {"dept": dept, "region": region_dir(dept), "status": "failed",
                  "failed_step": f"driver: {type(exc).__name__}",
                  "finished": datetime.now().strftime("%Y-%m-%d %H:%M")}
        mine[dept] = st
        save_status(mine, my_path)                    # only ever our own départements
        done = sum(1 for d in depts[:i] if mine.get(d, {}).get("status") == "ok")
        rate = (time.time() - t0) / i
        log.info(f"[{dept}] {st['status']}  rows={st.get('rows', '?')} "
                 f"tel={st.get('phones', '?')} mail={st.get('emails', '?')} "
                 f"joign={st.get('joignables', '?')}  |  {done}/{i} ok, "
                 f"~{rate * (len(depts) - i) / 60:.0f} min left"
                 + (f"  [worker {args.worker}]" if args.worker else ""))

    log.info("=" * 70)
    write_recap()
    print_status()


if __name__ == "__main__":
    main()
