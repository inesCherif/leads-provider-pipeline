"""
M6-S14 — Maha's call sheet for éleveurs: few columns, phones first, nothing she already has
==========================================================================================
The full file (m6_s9) is the audit trail. A phone operator needs the
opposite: one row per farm she can dial, every column meaningful, no row
she already received in another file.

Two tabs (Ines 2026-09-09, same layout as the gîtes sheet):
  Sûr       a dialled phone from a measured or owner-declared source
            (Agence Bio, OpenStreetMap, Pages Jaunes, bienvenue-à-la-ferme,
            the farm's own site, two independent sources agreeing), and the
            établissement active in SIRENE today.
  Probable  a phone known only from the provider file (73 % measured), or a
            provider row the registry could not confirm — labelled so.
Dropped, counted in the log: no phone at all; already sent to Maha (SIRET or
phone in any of the four files she holds); pet trades (dog / cat breeders);
public bodies (lycées agricoles, communes, associations); liquidation named
in the registry; a second row carrying a phone already on the sheet (one
farm = one call).

Gate (reads the xlsx BACK): T1 row count, T2 phone shape and no surtaxé,
T3 provenance filled, T4 SIRET unique, T5 no liquidation / closed row, T6
SIRET/CP as text, T7 e-mail shape / none invalide / no aggregator or mairie
domain, T8 e-mail statut iff e-mail, T9 no corporate domain on > 2 rows,
T10 no pet / public row, T11 Statut SIRENE filled, T14 type filled,
T16 NO SIRET AND NO PHONE PRESENT IN A FILE ALREADY SENT, T17 no provider-only
phone in the Sûr tab, T18 no phone twice on the sheet.

Usage:
    python scripts/m6_s14_export_teleop.py --departement 63 --version v1 --out-version v1
"""

import argparse
import csv
import logging
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))
if hasattr(sys.stdout, "reconfigure"):                 # Windows console is cp1252; labels carry accents
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from m1_s8_export import ILLEGAL_XML                                   # noqa: E402
from m2lib_contact import is_surtaxe, plausible_fr_number, FREE_MAIL   # noqa: E402
from m3ag_s12_check import PHONE_RE, EMAIL_RE, MAIRIE_RE, EXTRA_FREE_MAIL  # noqa: E402
from m6_lib import (CHECK_DIR, OUT_DIR, INHERITED_DIR, read_csv, is_aggregator,  # noqa: E402
                    PHONE_SOURCE_FR, load_sent, phone_digits)

COLUMNS = ["Entreprise", "Type d'élevage", "Bio", "Contact", "Téléphone", "Origine du téléphone",
           "E-mail", "Statut e-mail", "Adresse", "Code postal", "Ville",
           "SIRET", "Statut SIRENE", "Niveau"]

LISEZ_MOI = [
    ("Fichier", "Éleveurs — liste d'appel, département {dept}, générée le {today} (version {out})."),
    ("Onglet Sûr", "Une ligne par exploitation avec un téléphone issu d'une source identifiée (déclaré par "
                   "l'exploitant à l'Agence Bio, OpenStreetMap, Pages Jaunes, Bienvenue à la ferme, le site de "
                   "l'exploitation, ou deux sources indépendantes concordantes) et un établissement actif au "
                   "registre SIRENE à la date du fichier. C'est la liste à appeler en premier."),
    ("Onglet Probable", "Exploitations dont le téléphone ne vient que du fichier fournisseur (exact dans 73 % des cas "
                        "mesurés), ou que le registre n'a pas permis de confirmer. À appeler ensuite ; vérifier "
                        "l'identité en début d'appel."),
    ("Pas de doublon", "Aucune ligne de ce fichier ne figure dans les fichiers déjà envoyés (agriculteurs 63/03, "
                       "hébergements 63/03) : les SIRET et les numéros déjà transmis ont été retirés."),
    ("Exclus", "Sans téléphone ; déjà envoyés ; élevages d'animaux de compagnie (chiens, chats) ; organismes "
               "publics (lycées agricoles, communes, associations) ; liquidateur nommé au registre."),
    ("Entreprise", "Enseigne quand le registre en a une, sinon raison sociale (GAEC, EARL, nom de l'exploitant)."),
    ("Type d'élevage", "D'après le code d'activité du registre : bovins lait, bovins viande, ovins-caprins, "
                       "porcins, volailles, équins, autres animaux, polyculture-élevage."),
    ("Bio", "oui = exploitation certifiée en agriculture biologique (Agence Bio)."),
    ("Contact", "Gérant / exploitant inscrit au registre. Vide = inconnu, demander le responsable."),
    ("Téléphone", "Numéro à composer, format 0X XX XX XX XX. Jamais de numéro surtaxé. Un numéro n'apparaît "
                  "qu'une fois dans le fichier."),
    ("Origine du téléphone", "La source du numéro ; « confirmé par fichier fournisseur » = le fichier fournisseur "
                             "donne le même numéro."),
    ("E-mail / Statut", "vérifié = le serveur de messagerie a confirmé la boîte ; déclaré par l'exploitant = donné "
                        "par l'exploitation à l'Agence Bio ; non vérifié = trouvé sur le web, non confirmé. "
                        "Les adresses prouvées fausses ne sont pas dans le fichier."),
    ("SIRET / Statut SIRENE", "Identifiant de l'établissement (texte). actif = établissement actif au registre à la "
                              "date du fichier. non retrouvé = ligne issue du fichier fournisseur sans SIRET "
                              "vérifiable, onglet Probable."),
    ("Niveau", "Sûr ou Probable, la règle ci-dessus."),
]

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s %(levelname)-7s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("m6_s14")


def clean(v) -> str:
    return ILLEGAL_XML.sub("", str(v if v is not None else "")).strip()


def phone_source_fr(src: str) -> str:
    base = src.split("(")[0]
    if base == "corrobore":
        inside = src[src.index("(") + 1:].rstrip(")") if "(" in src else ""
        inside = inside.replace("fichier_fournisseur", "fichier fournisseur").replace("db_piste", "recherche web")
        return "2 sources indépendantes concordantes" + (f" ({inside})" if inside else "")
    return PHONE_SOURCE_FR.get(base, src)


def email_status_fr(statut: str) -> str:
    return {"valide": "vérifié", "risque": "risqué (le domaine accepte tout)", "declare": "déclaré par l'exploitant",
            "non verifie": "non vérifié"}.get(statut, "non vérifié")


def build(dept: str, version: str) -> tuple[list[dict], Counter]:
    src = OUT_DIR / f"eleveurs_{dept}_{version}.csv"
    if not src.exists():
        sys.exit(f"{src} missing — run m6_s9_export.py --version {version} first.")
    rows_in = read_csv(src, delim=",")
    stats = Counter(rows_in=len(rows_in))
    # Sûr before Probable, sièges first, so the phone-level dedup keeps the best row
    rows_in.sort(key=lambda r: (0 if r["telephone_final"] else 1, 0 if r["population_source"] == "registre" else 1,
                                0 if r.get("est_siege") == "Oui" else 1, r["codePostal"], r["raisonSociale"]))
    kept: dict[str, dict] = {}
    phones_used: set = set()
    for r in rows_in:
        provider_only = False
        tel, tel_src = r["telephone_final"], r["source_telephone"]
        if not tel and r["provider_phone"]:
            tel, tel_src, provider_only = r["provider_phone"], "provider", True
        if not tel:
            stats["dropped: no phone"] += 1
            continue
        if r["deja_envoye"]:
            stats[f"dropped: already sent to Maha ({r['deja_envoye']})"] += 1
            continue
        if r["flag_animaux_compagnie"] == "1":
            stats["dropped: pet trade (dogs / cats)"] += 1
            continue
        if r["flag_public"] == "1":
            stats["dropped: public body"] += 1
            continue
        if r["procedure_collective"]:
            stats["dropped: liquidation named in the registry"] += 1
            continue
        if not clean(r["raisonSociale"]):
            if clean(r["gerant"]):
                r["raisonSociale"] = r["gerant"]          # a sole trader IS the person
                stats["entreprise = contact name (no company name in the source)"] += 1
            else:
                stats["dropped: no company name and no contact name"] += 1
                continue
        if is_surtaxe(tel) or not plausible_fr_number(tel):
            stats["dropped: surtaxé / implausible number"] += 1
            continue
        d = phone_digits(tel)
        if d in phones_used:
            stats["merged: phone already on the sheet (one farm = one call)"] += 1
            continue
        key = r["siret"] if len(r["siret"]) == 14 else (f"X{r['siren']}" if r["siren"] else f"P{d}")
        if key in kept:
            stats["merged: duplicate SIRET"] += 1
            continue
        sirene_ok = r["statut_sirene"].startswith("actif")
        niveau = "Sûr" if (sirene_ok and not provider_only) else "Probable"
        if r["population_source"] == "provider":
            contact = clean(r["gerant"])
        else:
            contact = clean(r["gerant"]) if (r.get("nom") or "").strip() else (f"(gérée par {clean(r['gerant'])})" if r["gerant"] else "")
        origin = phone_source_fr(tel_src)
        if r["telephone_confirme_par"] and not provider_only:
            origin += f" — confirmé par {r['telephone_confirme_par']}"
        phones_used.add(d)
        kept[key] = {
            "Entreprise": clean(r["raisonSociale"]), "Type d'élevage": r["type_elevage"] or "non précisé",
            "Bio": "oui" if r["bio"] == "1" else "",
            "Contact": contact, "Téléphone": tel, "Origine du téléphone": origin,
            "E-mail": clean(r["email_final"]),
            "Statut e-mail": email_status_fr(r["email_statut"]) if r["email_final"] else "",
            "Adresse": clean(r["adresse"]), "Code postal": clean(r["codePostal"]), "Ville": clean(r["ville"]),
            "SIRET": r["siret"], "Statut SIRENE": r["statut_sirene"], "Niveau": niveau,
        }
        stats[f"kept: {niveau}"] += 1
    out = sorted(kept.values(), key=lambda x: (0 if x["Niveau"] == "Sûr" else 1, x["Code postal"], x["Entreprise"]))
    stats["rows_out"] = len(out)
    return out, stats


def write(rows: list[dict], dept: str, out_version: str) -> Path:
    csv_path = OUT_DIR / f"eleveurs_{dept}_teleop_{out_version}.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, delimiter=";")
        w.writeheader()
        w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    xlsx_path = OUT_DIR / f"eleveurs_{dept}_teleop_{out_version}.xlsx"
    wb = Workbook()
    widths = {"Entreprise": 34, "Type d'élevage": 18, "Bio": 5, "Contact": 24, "Téléphone": 16,
              "Origine du téléphone": 40, "E-mail": 32, "Statut e-mail": 24, "Adresse": 30,
              "Code postal": 11, "Ville": 22, "SIRET": 16, "Statut SIRENE": 16, "Niveau": 9}
    first = True
    for niveau in ("Sûr", "Probable"):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = niveau
        ws.append(COLUMNS)
        for r in rows:
            if r["Niveau"] == niveau:
                ws.append([str(r[c]) if c in ("SIRET", "Code postal") else r[c] for c in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, c in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths[c]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    lm = wb.create_sheet("Lisez-moi")
    lm.column_dimensions["A"].width = 26
    lm.column_dimensions["B"].width = 115
    for k, v in LISEZ_MOI:
        lm.append([k, v.format(dept=dept, today=date.today().strftime("%d/%m/%Y"), out=out_version)])
        lm.cell(row=lm.max_row, column=1).font = Font(bold=True)
        lm.cell(row=lm.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(xlsx_path)
    return xlsx_path


def gate(xlsx_path: Path, dept: str, version: str, expected: int) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True)
    rows = []
    for name in ("Sûr", "Probable"):
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        for vals in it:
            rows.append(dict(zip(header, vals)))
    verified = {}
    for p in (INHERITED_DIR / "verified_emails.csv", CHECK_DIR / "verified_emails.csv"):
        for r in read_csv(p):
            verified[r["email"].lower()] = r["verdict"]
    db_invalid = {r["email"].lower() for r in read_csv(CHECK_DIR / "db_claims.csv")
                  if r.get("kind") == "email" and r.get("verdict") in ("invalid", "invalide", "malformed")}
    full_rows = read_csv(OUT_DIR / f"eleveurs_{dept}_{version}.csv", delim=",")
    full = {r["siret"]: r for r in full_rows if len(r["siret"]) == 14}
    sent_sirets, sent_phones = load_sent()
    fails = []

    def check(cond, label):
        print(f"  {'ok  ' if cond else 'FAIL'}  {label}")
        if not cond:
            fails.append(label)

    g = lambda r, k: "" if r.get(k) is None else str(r[k])
    print(f"{xlsx_path.name}: {len(rows)} rows ({Counter(g(r,'Niveau') for r in rows)}), {len(COLUMNS)} columns")
    check(len(rows) == expected, f"T1 row count matches the build ({len(rows)} vs {expected})")
    check(all(g(r, "Téléphone") and PHONE_RE.match(g(r, "Téléphone")) for r in rows), "T2 every row has a phone in 0X XX XX XX XX shape")
    check(not any(is_surtaxe(g(r, "Téléphone")) or not plausible_fr_number(g(r, "Téléphone")) for r in rows), "T2 no surtaxé / implausible number")
    check(all(g(r, "Origine du téléphone") for r in rows), "T3 phone provenance always filled")
    sirets = [g(r, "SIRET") for r in rows if g(r, "SIRET")]
    check(len(sirets) == len(set(sirets)), f"T4 SIRET unique ({len(sirets) - len(set(sirets))} repeats)")
    dead = [s for s in sirets if full.get(s, {}).get("procedure_collective") or full.get(s, {}).get("statut_sirene", "").startswith("ferm")]
    check(not dead, f"T5 no liquidation / closed SIRET ({len(dead)})")
    check(all(isinstance(r.get("Code postal"), str) and len(g(r, "Code postal")) == 5 for r in rows)
          and all(isinstance(r.get("SIRET"), str) for r in rows if r.get("SIRET") not in (None, "")),
          "T6 SIRET and CP stored as text, CP 5 chars")
    mails = [g(r, "E-mail").lower() for r in rows if g(r, "E-mail")]
    check(all(EMAIL_RE.match(m) for m in mails), "T7 e-mail shape")
    bad_inv = [m for m in mails if verified.get(m) == "invalide" or m in db_invalid]
    check(not bad_inv, f"T7 no e-mail verified invalide ({len(bad_inv)})")
    bad_dom = [m for m in mails if (is_aggregator(m.rpartition("@")[2]) and m.rpartition("@")[2] != "gmail.com") or MAIRIE_RE.search(m.rpartition("@")[2])]
    check(not bad_dom, f"T7 no aggregator / mairie mailbox ({len(bad_dom)}) {bad_dom[:3]}")
    check(all(bool(g(r, "E-mail")) == bool(g(r, "Statut e-mail")) for r in rows), "T8 e-mail statut filled iff e-mail filled")
    dom = Counter(m.rpartition("@")[2] for m in mails if m.rpartition("@")[2] not in FREE_MAIL | EXTRA_FREE_MAIL)
    shared = [d for d, n in dom.items() if n > 2]
    check(not shared, f"T9 no corporate domain on > 2 rows ({shared[:3]})")
    bad_flag = [s for s in sirets if full.get(s, {}).get("flag_public") == "1" or full.get(s, {}).get("flag_animaux_compagnie") == "1"]
    check(not bad_flag, f"T10 no public body / pet-trade row ({len(bad_flag)})")
    check(all(g(r, "Statut SIRENE") for r in rows), "T11 Statut SIRENE always filled")
    check(all(g(r, "Type d'élevage") for r in rows), "T14 type d'élevage filled")
    check(all(g(r, "Entreprise").strip() for r in rows), f"T15 Entreprise never empty ({sum(1 for r in rows if not g(r, 'Entreprise').strip())})")
    dup_sent_s = [s for s in sirets if s in sent_sirets]
    dup_sent_p = [g(r, "Téléphone") for r in rows if phone_digits(g(r, "Téléphone")) in sent_phones]
    check(not dup_sent_s and not dup_sent_p,
          f"T16 no SIRET / phone already in a file sent to Maha ({len(dup_sent_s)} SIRET, {len(dup_sent_p)} phones)")
    prov_sur = [r for r in rows if g(r, "Niveau") == "Sûr" and g(r, "Origine du téléphone").startswith("fichier fournisseur")]
    check(not prov_sur, f"T17 no provider-only phone in the Sûr tab ({len(prov_sur)})")
    ph = Counter(phone_digits(g(r, "Téléphone")) for r in rows)
    check(not [p for p, n in ph.items() if n > 1], f"T18 no phone twice on the sheet ({sum(1 for n in ph.values() if n > 1)})")
    print(f"  info  Contact filled: {sum(1 for r in rows if g(r, 'Contact'))}/{len(rows)}; e-mail: {len(mails)} "
          f"({Counter(g(r, 'Statut e-mail') for r in rows if g(r, 'E-mail')).most_common()}); bio: {sum(1 for r in rows if g(r, 'Bio'))}; "
          f"types: {Counter(g(r, chr(84)+'ype d'+chr(39)+'élevage') for r in rows).most_common(8)}")
    print(f"{len(fails)} failed")
    return 1 if fails else 0


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--departement", default="63")
    ap.add_argument("--version", default="v1", help="m6_s9 version to read")
    ap.add_argument("--out-version", default="v1")
    args = ap.parse_args()

    rows, stats = build(args.departement, args.version)
    xlsx_path = write(rows, args.departement, args.out_version)
    log.info("─" * 62)
    log.info(f"[{args.departement}] {stats['rows_in']} rows in -> {stats['rows_out']} rows out -> {xlsx_path.name} + .csv")
    for k, v in stats.most_common():
        if k not in ("rows_in", "rows_out"):
            log.info(f"  {k:<64} {v:>5}")
    log.info(f"  phone origin: {dict(Counter(r['Origine du téléphone'].split(' —')[0] for r in rows))}")
    log.info("reading the xlsx back:")
    sys.exit(gate(xlsx_path, args.departement, args.version, len(rows)))


if __name__ == "__main__":
    main()
