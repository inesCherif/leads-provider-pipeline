"""
FRANCE-VERIFY — read every delivered file BACK and check the whole country
=========================================================================
The per-département gates (m7_s12, m6_s12) prove one file at a time. This
proves the DELIVERY: that the 96 files exist, that none of them breaks the
principle, that every row sits in the right département, and that the recap
totals are the sum of the files rather than of the logs.

Everything here reads the shipped .xlsx. Nothing trusts a log line — the rule
this project learned when `collected=` hid 3,402 lost rows.

    python scripts/france_verify.py
    python scripts/france_verify.py --strict     # warnings fail too
"""

import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))
import re                                                       # noqa: E402
from m7_lib import (EXCLUDED_NAF, EXCLUDED_RE, host_hits,       # noqa: E402
                    load_principle_rescue)
from france_lib import METRO_DEPARTEMENTS, REGION_OF, in_dept   # noqa: E402
import france_run as fr                                         # noqa: E402

NAME_FIELDS = ("raisonSociale", "denomination_legale", "enseigne")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--strict", action="store_true")
    args = ap.parse_args()

    from openpyxl import load_workbook

    rescued = load_principle_rescue()
    fails, warns = [], []

    def check(cond, label, warn=False):
        if cond:
            print(f"  ok    {label}")
        else:
            (warns if warn else fails).append(label)
            print(f"  {'WARN' if warn else 'FAIL'}  {label}")

    # ---- which files are actually on disk, per région ----
    found, missing = {}, []
    for dept in METRO_DEPARTEMENTS:
        reg = fr.OUT_ROOT / REGION_OF[dept]
        hits = sorted(reg.glob(f"producteurs_{dept}_v*.xlsx")) if reg.exists() else []
        if hits:
            found[dept] = hits[-1]          # v2 sorts after v1: the delivered one
        else:
            missing.append(dept)
    print(f"\nFRANCE VERIFY — {len(found)}/{len(METRO_DEPARTEMENTS)} départements on disk\n")
    check(not missing, f"F1 every département delivered ({len(missing)} missing) "
                       f"{missing[:15]}")

    tot = Counter()
    bad_naf, bad_name, bad_host, bad_dept, empty = [], [], [], [], []
    siret_seen: dict[str, str] = {}
    dup_siret: list[tuple[str, str, str]] = []
    segs = Counter()
    per_region = defaultdict(Counter)

    for dept, path in found.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb["Producteurs"]
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                empty.append(dept)
                continue
            ix = {n: i for i, n in enumerate(header)}

            def g(r, k):
                i = ix.get(k)
                v = r[i] if i is not None and i < len(r) else None
                return "" if v is None else str(v)

            n = 0
            for r in it:
                if not r or all(v in (None, "") for v in r):
                    continue
                n += 1
                naf = g(r, "codeNAF")
                if naf in EXCLUDED_NAF:
                    bad_naf.append((dept, g(r, "siret"), naf))
                # A reviewed surname collision is kept on purpose (Vigneron,
                # Brasseur and Pinot are ordinary French names), so this must
                # apply the same rescue list the gate does or it reports the
                # decision itself as a defect.
                if re.sub(r"\D", "", g(r, "siret")) not in rescued \
                        and EXCLUDED_RE.search(" ".join(g(r, k) for k in NAME_FIELDS)):
                    bad_name.append((dept, g(r, "raisonSociale")))
                site, mail = g(r, "website_final"), g(r, "email_final")
                if (site and host_hits(site.split("//")[-1].split("/")[0])) or \
                   (mail and host_hits(mail.rpartition("@")[2])):
                    bad_host.append((dept, g(r, "raisonSociale"), site or mail))
                cp = g(r, "codePostal")
                if cp and not in_dept(cp, dept):
                    bad_dept.append((dept, cp))
                s = g(r, "siret")
                if s and len(s) == 14:
                    if s in siret_seen and siret_seen[s] != dept:
                        dup_siret.append((s, siret_seen[s], dept))
                    siret_seen[s] = dept
                for t in g(r, "sous_segment").split("|"):
                    if t:
                        segs[t] += 1
                tot["phones"] += 1 if g(r, "telephone_final") else 0
                tot["emails"] += 1 if g(r, "email_final") else 0
                tot["provider_phones"] += 1 if g(r, "provider_phone") else 0
                if g(r, "telephone_final") or g(r, "email_final"):
                    tot["joignables"] += 1
                if g(r, "telephone_final") or g(r, "email_final") or g(r, "provider_phone"):
                    tot["joignables_total"] += 1
                if "eleveurs" in g(r, "population_source"):
                    tot["eleveurs"] += 1
                else:
                    tot["producteurs"] += 1
            tot["rows"] += n
            per_region[REGION_OF[dept]]["rows"] += n
            per_region[REGION_OF[dept]]["depts"] += 1
            if n == 0:
                empty.append(dept)
        finally:
            wb.close()

    print()
    check(not empty, f"F2 no empty département file ({len(empty)}) {empty[:10]}")
    check(not bad_naf, f"F3 principle: no excluded NAF anywhere ({len(bad_naf)}) {bad_naf[:3]}")
    check(not bad_name, f"F4 principle: no excluded word in a business name ({len(bad_name)}) {bad_name[:3]}")
    check(not bad_host, f"F5 principle: no wine / beer / pork domain shipped ({len(bad_host)}) {bad_host[:2]}")
    check(not bad_dept, f"F6 every postcode sits in its own département ({len(bad_dept)}) {bad_dept[:5]}")
    check(not dup_siret, f"F7 a SIRET appears in one département only ({len(dup_siret)}) {dup_siret[:3]}",
          warn=True)

    print(f"\n{'région':<28}{'depts':>7}{'lignes':>10}")
    print("-" * 45)
    for reg in sorted(per_region):
        print(f"{reg:<28}{per_region[reg]['depts']:>7}{per_region[reg]['rows']:>10,}".replace(",", " "))
    print("-" * 45)
    n = tot["rows"]
    print(f"{'TOTAL':<28}{len(found):>7}{n:>10,}".replace(",", " "))
    print(f"\n  producteurs            {tot['producteurs']:>8,}".replace(",", " "))
    print(f"  éleveurs               {tot['eleveurs']:>8,}".replace(",", " "))
    print(f"  téléphones (mesurés)   {tot['phones']:>8,}  ({tot['phones']/n:.1%})".replace(",", " ") if n else "")
    print(f"  e-mails                {tot['emails']:>8,}  ({tot['emails']/n:.1%})".replace(",", " ") if n else "")
    print(f"  tél. fichier client    {tot['provider_phones']:>8,}  ({tot['provider_phones']/n:.1%})".replace(",", " ") if n else "")
    print(f"  joignables (mesurés)   {tot['joignables']:>8,}  ({tot['joignables']/n:.1%})".replace(",", " ") if n else "")
    print(f"  joignables + fichier   {tot['joignables_total']:>8,}  ({tot['joignables_total']/n:.1%})".replace(",", " ") if n else "")
    print(f"\n  top sous-segments: {dict(segs.most_common(10))}")

    print(f"\n{len(fails)} failed, {len(warns)} warning(s)")
    sys.exit(1 if fails or (args.strict and warns) else 0)


if __name__ == "__main__":
    main()
