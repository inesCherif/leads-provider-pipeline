"""
M5-S14 — Maha's call sheet: few columns, only what we stand behind, with the proof of life
=========================================================================================
The full file (m5_s9) is the audit trail. A phone operator needs the
opposite: one row per business she can dial, every column meaningful, the
reason we believe the business is open written in plain French.

Two tabs (Ines 2026-09-09):
  Sûr       a dialled phone from a measured or owner-declared source, SIRENE
            actif, no BODACC death or procedure, ≥ 2 independent proofs of
            activity (offre d'emploi, dépôt des comptes, fiche office de
            tourisme, classement Atout France, fiche annuaire, immatriculation
            récente, site actif).
  Probable  the rest worth a call: 1 proof, or SIRENE not found, or a phone
            known only from the provider file (73 % measured) — labelled so.
Dropped, counted in the log: no phone at all; SIRENE fermé; BODACC radié /
liquidation / fonds cédé; redressement / sauvegarde (a business in a
procedure does not buy energy works — Ines can reverse it); public
operators (phase 2); hotels.

Gate (reads the xlsx BACK): T1 row count, T2 phone shape and no surtaxé,
T3 provenance filled, T4 SIRET unique, T5 no closed SIRET, T6 SIRET/CP as
text, T7 e-mail shape / none invalide / no aggregator or mairie domain,
T8 e-mail statut iff e-mail, T9 no corporate domain on > 2 rows,
T10 no public / hors-secteur row, T11 Statut SIRENE filled, T12 no BODACC-dead
SIREN, T13 every Sûr row has ≥ 2 proofs, T14 type filled, T15 Dernière
annonce dated ≤ today with an https URL.

Usage:
    python scripts/m5_s14_export_teleop.py --departement 63 --version v0 --out-version v0
"""

import argparse
import csv
import logging
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))
if hasattr(sys.stdout, "reconfigure"):                 # Windows console is cp1252; labels carry ≥ and accents
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from m1_s8_export import ILLEGAL_XML                                   # noqa: E402
from m2lib_contact import is_surtaxe, plausible_fr_number, FREE_MAIL   # noqa: E402
from m3ag_s12_check import PHONE_RE, EMAIL_RE, MAIRIE_RE, EXTRA_FREE_MAIL, load_xlsx  # noqa: E402
from m5_lib import CHECK_DIR, OUT_DIR, read_csv, is_aggregator, PHONE_SOURCE_FR  # noqa: E402

COLUMNS = ["Entreprise", "Type d'hébergement", "Contact", "Téléphone", "Origine du téléphone",
           "E-mail", "Statut e-mail", "Adresse", "Code postal", "Ville",
           "Capacité / emplacements", "Classement / labels", "Site web",
           "Preuves d'activité", "Dernière annonce", "SIRET", "Statut SIRENE", "Niveau"]
OUT_TYPES_DROP = {"hôtel"}

LISEZ_MOI = [
    ("Fichier", "Gîtes et campings — liste d'appel, département {dept}, générée le {today} (version {out})."),
    ("Onglet Sûr", "Une ligne par établissement avec un téléphone issu d'une source identifiée, actif au registre "
                   "SIRENE, sans radiation ni procédure au BODACC, et au moins DEUX preuves d'activité indépendantes "
                   "et datées. C'est la liste à appeler en premier."),
    ("Onglet Probable", "Établissements avec une seule preuve d'activité, ou dont le SIRET n'a pas été retrouvé, ou "
                        "dont le téléphone ne vient que du fichier fournisseur (exact dans 73 % des cas mesurés). "
                        "À appeler ensuite ; vérifier l'identité en début d'appel."),
    ("Exclus", "Sans téléphone ; fermés au registre ; radiés, liquidés ou fonds cédé au BODACC ; en redressement "
               "ou sauvegarde ; opérateurs publics (mairies, communautés de communes — phase 2) ; hôtels."),
    ("Entreprise", "Enseigne quand le registre en a une, sinon raison sociale."),
    ("Type d'hébergement", "camping, gîte, meublé, chambres d'hôtes, résidence, village vacances ; 'non typé' = "
                           "personne physique louant un meublé, type inconnu."),
    ("Contact", "Gérant / président / exploitant inscrit au registre. Vide = inconnu, demander le responsable."),
    ("Téléphone", "Numéro à composer, format 0X XX XX XX XX. Jamais de numéro surtaxé."),
    ("Origine du téléphone", "OpenStreetMap (concordance 96 % mesurée), Pages Jaunes (83-89 %), le site de "
                             "l'établissement, deux sources indépendantes concordantes, une offre d'emploi France "
                             "Travail (déclaré par l'établissement), ou le fichier fournisseur (Probable uniquement)."),
    ("E-mail / Statut", "vérifié = le serveur de messagerie a confirmé la boîte ; déclaré = donné par l'établissement "
                        "(offre d'emploi, fichier fournisseur) ; non vérifié = trouvé sur le web, non confirmé. "
                        "Les adresses prouvées fausses ne sont pas dans le fichier."),
    ("Capacité / emplacements", "Capacité d'accueil (personnes) et nombre d'emplacements, d'après le classement "
                                "Atout France ou OpenStreetMap. Un indicateur de taille pour l'argumentaire énergie."),
    ("Classement / labels", "Étoiles Atout France ; labels de la fiche office de tourisme (Gîtes de France, épis, "
                            "Clévacances, Qualité Tourisme…)."),
    ("Preuves d'activité", "Chaque élément est une trace publique DATÉE et indépendante : offre d'emploi France "
                           "Travail, dépôt des comptes / modification / création au BODACC, fiche office de tourisme "
                           "mise à jour, classement Atout France en cours de validité, fiche OpenStreetMap / Pages "
                           "Jaunes, immatriculation récente, site web actif."),
    ("Dernière annonce", "La plus récente annonce publique (BODACC ou France Travail) : type · date · lien."),
    ("SIRET / Statut SIRENE", "Identifiant de l'établissement (texte). actif = établissement actif au registre à la "
                              "date du fichier. SIRET inconnu / non retrouvé = pas de vérification possible, ligne en "
                              "Probable."),
    ("Niveau", "Sûr ou Probable, la règle ci-dessus."),
]

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s %(levelname)-7s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("m5_s14")
DEAD = {"radié", "liquidation judiciaire", "fonds cédé"}
FLAGGED = {"redressement judiciaire", "sauvegarde", "plan de redressement (en cours)", "procédure collective (autre)"}


def clean(v) -> str:
    return ILLEGAL_XML.sub("", str(v if v is not None else "")).strip()


def phone_source_fr(src: str) -> str:
    base = src.split("(")[0]
    if base == "corrobore":
        return "2 sources indépendantes concordantes" + (" " + src[src.index("("):] if "(" in src else "")
    return PHONE_SOURCE_FR.get(base, src)


def email_status_fr(statut: str) -> str:
    return {"valide": "vérifié", "risque": "risqué (le domaine accepte tout)", "declare": "déclaré par l'établissement",
            "non verifie": "non vérifié"}.get(statut, "non vérifié")


def build(dept: str, version: str) -> tuple[list[dict], Counter]:
    src = OUT_DIR / f"hebergement_{dept}_{version}.csv"
    if not src.exists():
        sys.exit(f"{src} missing — run m5_s9_export.py --version {version} first.")
    rows_in = read_csv(src, delim=",")
    stats = Counter(rows_in=len(rows_in))
    kept: dict[str, dict] = {}
    for r in rows_in:
        provider_only = ""
        tel, tel_src = r["telephone_final"], r["source_telephone"]
        if not tel:
            m = re.match(r"^(0\d(?: \d\d){4}) \(fichier fournisseur\)", r["telephone_piste"] or "")
            if m:
                tel, tel_src, provider_only = m.group(1), "provider", "1"
        if not tel:
            stats["dropped: no phone"] += 1
            continue
        if r["flag_public"] == "1":
            stats["dropped: public operator (phase 2)"] += 1
            continue
        if r["type_final"] in OUT_TYPES_DROP:
            stats["dropped: hotel"] += 1
            continue
        if r["mort"] == "1":
            stats["dropped: dead (SIRENE fermé / BODACC radié-liquidé-cédé)"] += 1
            continue
        if r["bodacc_statut"] in FLAGGED:
            stats[f"dropped: BODACC {r['bodacc_statut']}"] += 1
            continue
        score = int(r["score_activite"] or 0)
        sirene_ok = r["statut_sirene"].startswith("actif")
        if sirene_ok and score >= 2 and not provider_only and not r["alerte"]:
            niveau = "Sûr"
        elif score >= 1 or provider_only:
            niveau = "Probable"
        else:
            stats["dropped: no proof of activity"] += 1
            continue
        key = r["siret"] if len(r["siret"]) == 14 else f"X{r['siren']}"
        if key in kept:
            stats["merged: duplicate SIRET"] += 1
            continue
        cap = " / ".join(x for x in (r["capacite"] and f"{r['capacite']} pers.", r["emplacements"] and f"{r['emplacements']} empl.") if x)
        labels = " ; ".join(x for x in (r["classement"], r["labels"]) if x)
        origin = phone_source_fr(tel_src)
        if r["telephone_confirme_par"] and not provider_only:
            origin += f" — confirmé par {r['telephone_confirme_par'].split(' (')[0] if ' (' not in r['telephone_confirme_par'] else r['telephone_confirme_par'].split('(')[1].rstrip(')')}"
        kept[key] = {
            "Entreprise": clean(r["raisonSociale"]), "Type d'hébergement": r["type_final"] or "non typé",
            "Contact": clean(r["gerant"]), "Téléphone": tel, "Origine du téléphone": origin,
            "E-mail": clean(r["email_final"]), "Statut e-mail": email_status_fr(r["email_statut"]) if r["email_final"] else "",
            "Adresse": clean(r["adresse"]), "Code postal": clean(r["codePostal"]), "Ville": clean(r["ville"]),
            "Capacité / emplacements": cap, "Classement / labels": clean(labels)[:200], "Site web": clean(r["website_final"]),
            "Preuves d'activité": clean(r["preuves_activite"]), "Dernière annonce": clean(r["derniere_annonce"]),
            "SIRET": r["siret"], "Statut SIRENE": r["statut_sirene"], "Niveau": niveau,
        }
        stats[f"kept: {niveau}"] += 1
    out = sorted(kept.values(), key=lambda x: (0 if x["Niveau"] == "Sûr" else 1, x["Code postal"], x["Entreprise"]))
    stats["rows_out"] = len(out)
    return out, stats


def write(rows: list[dict], dept: str, out_version: str) -> Path:
    csv_path = OUT_DIR / f"hebergement_{dept}_teleop_{out_version}.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, delimiter=";")
        w.writeheader()
        w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    xlsx_path = OUT_DIR / f"hebergement_{dept}_teleop_{out_version}.xlsx"
    wb = Workbook()
    widths = {"Entreprise": 34, "Type d'hébergement": 16, "Contact": 24, "Téléphone": 16, "Origine du téléphone": 34,
              "E-mail": 32, "Statut e-mail": 22, "Adresse": 30, "Code postal": 11, "Ville": 22,
              "Capacité / emplacements": 18, "Classement / labels": 30, "Site web": 30,
              "Preuves d'activité": 60, "Dernière annonce": 50, "SIRET": 16, "Statut SIRENE": 16, "Niveau": 9}
    first = True
    for niveau in ("Sûr", "Probable"):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = niveau
        ws.append(COLUMNS)
        for r in rows:
            if r["Niveau"] == niveau:
                ws.append([str(r[c]) if c in ("SIRET", "Code postal") else r[c] for c in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, c in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths[c]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    lm = wb.create_sheet("Lisez-moi")
    lm.column_dimensions["A"].width = 26
    lm.column_dimensions["B"].width = 115
    for k, v in LISEZ_MOI:
        lm.append([k, v.format(dept=dept, today=date.today().strftime("%d/%m/%Y"), out=out_version)])
        lm.cell(row=lm.max_row, column=1).font = Font(bold=True)
        lm.cell(row=lm.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(xlsx_path)
    return xlsx_path


def gate(xlsx_path: Path, dept: str, version: str, expected: int) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True)
    rows = []
    for name in ("Sûr", "Probable"):
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        for vals in it:
            rows.append(dict(zip(header, vals)))
    sirene = {r["siret"]: r for r in read_csv(CHECK_DIR / f"sirene_etat_{dept}.csv")}
    bodacc = {r["siren"]: r for r in read_csv(CHECK_DIR / f"bodacc_status_{dept}.csv")}
    verified = {r["email"].lower(): r["verdict"] for r in read_csv(CHECK_DIR / "verified_emails.csv")}
    full = {r["siret"]: r for r in read_csv(OUT_DIR / f"hebergement_{dept}_{version}.csv", delim=",") if len(r["siret"]) == 14}
    fails = []

    def check(cond, label):
        print(f"  {'ok  ' if cond else 'FAIL'}  {label}")
        if not cond:
            fails.append(label)

    g = lambda r, k: "" if r.get(k) is None else str(r[k])
    print(f"{xlsx_path.name}: {len(rows)} rows ({Counter(g(r,'Niveau') for r in rows)}), {len(COLUMNS)} columns")
    check(len(rows) == expected, f"T1 row count matches the build ({len(rows)} vs {expected})")
    check(all(g(r, "Téléphone") and PHONE_RE.match(g(r, "Téléphone")) for r in rows), "T2 every row has a phone in 0X XX XX XX XX shape")
    check(not any(is_surtaxe(g(r, "Téléphone")) or not plausible_fr_number(g(r, "Téléphone")) for r in rows), "T2 no surtaxé / implausible number")
    check(all(g(r, "Origine du téléphone") for r in rows), "T3 phone provenance always filled")
    sirets = [g(r, "SIRET") for r in rows if g(r, "SIRET")]
    check(len(sirets) == len(set(sirets)), f"T4 SIRET unique ({len(sirets) - len(set(sirets))} repeats)")
    closed = [s for s in sirets if s in sirene and (sirene[s]["etat_ul"] == "C" or sirene[s]["etat_etab"] == "F")]
    check(not closed, f"T5 no SIRET closed in SIRENE ({len(closed)})")
    check(all(isinstance(r.get("SIRET"), str) and isinstance(r.get("Code postal"), str) and len(g(r, "Code postal")) == 5 for r in rows), "T6 SIRET and CP stored as text, CP 5 chars")
    mails = [g(r, "E-mail").lower() for r in rows if g(r, "E-mail")]
    check(all(EMAIL_RE.match(m) for m in mails), "T7 e-mail shape")
    check(not any(verified.get(m) == "invalide" for m in mails), "T7 no e-mail verified invalide")
    bad_dom = [m for m in mails if (is_aggregator(m.rpartition("@")[2]) and m.rpartition("@")[2] != "gmail.com") or MAIRIE_RE.search(m.rpartition("@")[2])]
    check(not bad_dom, f"T7 no aggregator / mairie mailbox ({len(bad_dom)}) {bad_dom[:3]}")
    check(all(bool(g(r, "E-mail")) == bool(g(r, "Statut e-mail")) for r in rows), "T8 e-mail statut filled iff e-mail filled")
    dom = Counter(m.rpartition("@")[2] for m in mails if m.rpartition("@")[2] not in FREE_MAIL | EXTRA_FREE_MAIL)
    shared = [d for d, n in dom.items() if n > 2]
    check(not shared, f"T9 no corporate domain on > 2 rows ({shared[:3]})")
    pub = [r for r in rows if full.get(g(r, "SIRET"), {}).get("flag_public") == "1" or g(r, "Type d'hébergement") in OUT_TYPES_DROP]
    check(not pub, f"T10 no public operator / hotel row ({len(pub)})")
    check(all(g(r, "Statut SIRENE") for r in rows), "T11 Statut SIRENE always filled")
    dead = [r for r in rows if bodacc.get(g(r, "SIRET")[:9], {}).get("statut") in DEAD | FLAGGED]
    check(not dead, f"T12 no BODACC dead / flagged SIREN ({len(dead)})")
    weak = [r for r in rows if g(r, "Niveau") == "Sûr" and g(r, "Preuves d'activité").count(" · ") + 1 < 2]
    check(not weak, f"T13 every Sûr row has ≥ 2 proofs of activity ({len(weak)})")
    check(all(g(r, "Type d'hébergement") for r in rows), "T14 type d'hébergement filled")
    today = date.today().strftime("%Y-%m-%d")
    bad_ann = []
    for r in rows:
        a = g(r, "Dernière annonce")
        if not a:
            continue
        m = re.search(r"· (\d\d)/(\d\d)/(\d{4}) · (\S+)", a)
        if not m or f"{m.group(3)}-{m.group(2)}-{m.group(1)}" > today or not m.group(4).startswith("https://"):
            bad_ann.append(a)
    check(not bad_ann, f"T15 Dernière annonce dated ≤ today with an https URL ({len(bad_ann)})")
    type_col = "Type d'hébergement"
    print(f"  info  Contact filled: {sum(1 for r in rows if g(r, 'Contact'))}/{len(rows)}; e-mail: {len(mails)} "
          f"({Counter(g(r, 'Statut e-mail') for r in rows if g(r, 'E-mail')).most_common()}); "
          f"types: {Counter(g(r, type_col) for r in rows).most_common(6)}")
    print(f"{len(fails)} failed")
    return 1 if fails else 0


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--departement", default="63")
    ap.add_argument("--version", default="v0", help="m5_s9 version to read")
    ap.add_argument("--out-version", default="v0")
    args = ap.parse_args()

    rows, stats = build(args.departement, args.version)
    xlsx_path = write(rows, args.departement, args.out_version)
    log.info("─" * 62)
    log.info(f"[{args.departement}] {stats['rows_in']} rows in -> {stats['rows_out']} rows out -> {xlsx_path.name} + .csv")
    for k, v in stats.most_common():
        if k not in ("rows_in", "rows_out"):
            log.info(f"  {k:<60} {v:>5}")
    log.info(f"  phone origin: {dict(Counter(r['Origine du téléphone'].split(' —')[0] for r in rows))}")
    log.info("reading the xlsx back:")
    sys.exit(gate(xlsx_path, args.departement, args.version, len(rows)))


if __name__ == "__main__":
    main()
