"""
maha_lib — the "never twice" rule for every call sheet sent to Maha
===================================================================
Ines's rule (2026-09-09): a file sent to Maha (téléopératrice) must never
contain a company she already received — in ANY earlier file, ANY sector,
ANY version. A V2 holds only the NEW companies.

Mechanism, deliberately dumb:
  * exports/maha_sent/  holds a copy of every xlsx actually sent (gitignored
    with the rest of exports/; `docs/maha_deliveries.md` is the tracked log).
  * `load_sent()` reads every xlsx there BACK (all sheets) and returns the
    SIRETs and the phone numbers they contain.
  * every *_s14_export_teleop.py drops a row whose SIRET or phone is in that
    set and re-checks it in its gate (T16) from the built xlsx.
  * `register(path)` copies a just-sent file into the folder and appends a
    line to the log. Run it right after sending, before the next rebuild —
    a rebuild of the same version overwrites the file in exports/<sector>/.

    python scripts/maha_lib.py --register exports/eleveurs/eleveurs_03_teleop_v2.xlsx
    python scripts/maha_lib.py --list
"""

from __future__ import annotations

import re
import shutil
import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
SENT_DIR = PROJECT_ROOT / "exports" / "maha_sent"
LOG_PATH = PROJECT_ROOT / "docs" / "maha_deliveries.md"


def phone_digits(s: str) -> str:
    d = re.sub(r"\D", "", s or "")
    if d.startswith("33") and len(d) == 11:
        d = "0" + d[2:]
    return d if len(d) == 10 else ""


def sent_files() -> list[Path]:
    return sorted(SENT_DIR.glob("*.xlsx")) if SENT_DIR.exists() else []


def load_sent() -> tuple[set, set]:
    """(sirets, phone_digits) present in every file Maha already received.
    Reads the xlsx BACK (every sheet). An empty folder is a hard stop: the
    exclusion is the whole point, a missing folder is a mistake, not zero."""
    from openpyxl import load_workbook
    files = sent_files()
    if not files:
        sys.exit(f"{SENT_DIR} holds no xlsx — copy the files Maha already received there first "
                 "(python scripts/maha_lib.py --register <file>).")
    sirets, phones = set(), set()
    for path in files:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            cols = {str(h or "").strip().lower(): i for i, h in enumerate(header)}
            i_siret = cols.get("siret")
            i_tel = next((i for h, i in cols.items() if h.startswith("t") and "phone" in h), None)
            for r in rows:
                if i_siret is not None and i_siret < len(r) and r[i_siret]:
                    s = re.sub(r"\D", "", str(r[i_siret]))
                    if len(s) == 14:
                        sirets.add(s)
                if i_tel is not None and i_tel < len(r) and r[i_tel]:
                    d = phone_digits(str(r[i_tel]))
                    if d:
                        phones.add(d)
        wb.close()
    return sirets, phones


def register(path: Path, note: str = "") -> Path:
    """Copy a sent file into SENT_DIR and log it. Same name = overwrite (a V1
    re-sent after a rebuild replaces the V1 copy)."""
    if not path.exists():
        sys.exit(f"{path} not found")
    SENT_DIR.mkdir(parents=True, exist_ok=True)
    dst = SENT_DIR / path.name
    shutil.copy2(path, dst)
    from openpyxl import load_workbook
    wb = load_workbook(dst, read_only=True)
    n = sum(max(0, ws.max_row - 1) for ws in wb.worksheets if ws.title != "Lisez-moi")
    wb.close()
    if not LOG_PATH.exists():
        LOG_PATH.write_text("# Files sent to Maha (téléopératrice) — the 'never twice' register\n\n"
                            "One line per file actually sent. `exports/maha_sent/` holds the copies; "
                            "every call sheet excludes their SIRETs and phones (maha_lib).\n\n"
                            "| sent | file | rows | note |\n|---|---|---|---|\n", encoding="utf-8")
    with LOG_PATH.open("a", encoding="utf-8") as fh:
        fh.write(f"| {date.today().isoformat()} | `{path.name}` | {n} | {note} |\n")
    print(f"registered {dst.name}: {n} rows; sent files now {len(sent_files())}")
    return dst


if __name__ == "__main__":
    if "--register" in sys.argv:
        i = sys.argv.index("--register")
        note = " ".join(sys.argv[i + 2:])
        register(Path(sys.argv[i + 1]), note)
    elif "--list" in sys.argv:
        s, p = load_sent()
        for f in sent_files():
            print(f"  {f.name}")
        print(f"{len(sent_files())} file(s): {len(s)} SIRET, {len(p)} phones excluded from the next sheets")
    else:
        print(__doc__)
